from io import BytesIO

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.core.config import get_settings


PARSE_URL = "/api/v1/internal/documents/parse"


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Printer Manual", level=1)
    document.add_paragraph("Turn the printer off before maintenance.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Code"
    table.cell(0, 1).text = "Meaning"
    table.cell(1, 0).text = "E01"
    table.cell(1, 1).text = "Paper jam"

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf_bytes(
    *,
    include_text: bool = True,
    encryption_password: str | None = None,
) -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)

    if include_text:
        font = DictionaryObject(
            {
                NameObject("/Type"): NameObject("/Font"),
                NameObject("/Subtype"): NameObject("/Type1"),
                NameObject("/BaseFont"): NameObject("/Helvetica"),
            }
        )
        font_reference = writer._add_object(font)
        page[NameObject("/Resources")] = DictionaryObject(
            {
                NameObject("/Font"): DictionaryObject(
                    {NameObject("/F1"): font_reference}
                )
            }
        )
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (Hello PDF) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)

    if encryption_password is not None:
        writer.encrypt(encryption_password)

    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def _encrypted_pdf_bytes() -> bytes:
    return _pdf_bytes(encryption_password="test-password")


def test_parse_requires_internal_token(client: TestClient) -> None:
    response = client.post(
        PARSE_URL,
        files={"file": ("manual.txt", b"hello", "text/plain")},
    )

    assert response.status_code == 401
    assert response.json()["code"] == "UNAUTHORIZED"


def test_parse_rejects_invalid_internal_token(client: TestClient) -> None:
    response = client.post(
        PARSE_URL,
        headers={"X-OfficeCare-Internal-Token": "wrong-token"},
        files={"file": ("manual.txt", b"hello", "text/plain")},
    )

    assert response.status_code == 401
    assert response.json()["code"] == "UNAUTHORIZED"


def test_parse_reports_missing_service_configuration(
    client: TestClient,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    # Empty environment value takes precedence over a developer's ignored .env file.
    monkeypatch.setenv("OFFICECARE_AI_SERVICE_TOKEN", "")
    get_settings.cache_clear()

    response = client.post(
        PARSE_URL,
        headers={"X-OfficeCare-Internal-Token": "anything"},
        files={"file": ("manual.txt", b"hello", "text/plain")},
    )

    assert response.status_code == 503
    assert response.json()["code"] == "SERVICE_NOT_CONFIGURED"


def test_parse_requires_file_part(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(PARSE_URL, headers=auth_headers)

    assert response.status_code == 422
    assert response.json()["code"] == "REQUEST_VALIDATION_ERROR"
    assert response.json()["details"]["errors"]


def test_parse_utf8_txt_preserves_paragraph_order(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    content = "第一段：安装说明。\n\n第二段：维护说明。".encode()

    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={"file": ("../printer-manual.TXT", content, "text/plain")},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["fileName"] == "printer-manual.TXT"
    assert body["fileType"] == "txt"
    assert body["contentType"] == "text/plain"
    assert body["fileSize"] == len(content)
    assert len(body["fileHash"]) == 64
    assert body["pageCount"] is None
    assert body["blockCount"] == 2
    assert [block["sequence"] for block in body["blocks"]] == [1, 2]
    assert [block["text"] for block in body["blocks"]] == [
        "第一段：安装说明。",
        "第二段：维护说明。",
    ]
    assert all(block["pageNumber"] is None for block in body["blocks"])
    assert body["warnings"] == []


def test_parse_gb18030_txt_adds_warning(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={
            "file": (
                "legacy.txt",
                "打印机维护说明".encode("gb18030"),
                "text/plain",
            )
        },
    )

    assert response.status_code == 200
    assert response.json()["blocks"][0]["text"] == "打印机维护说明"
    assert response.json()["warnings"] == [
        "The TXT file was decoded using GB18030 instead of UTF-8."
    ]


def test_parse_docx_preserves_heading_paragraph_and_table_order(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={
            "file": (
                "printer.docx",
                _docx_bytes(),
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["fileType"] == "docx"
    assert body["pageCount"] is None
    assert [block["blockType"] for block in body["blocks"]] == [
        "heading",
        "paragraph",
        "table",
    ]
    assert body["blocks"][0]["text"] == "Printer Manual"
    assert body["blocks"][2]["text"] == (
        "Code\tMeaning\nE01\tPaper jam"
    )
    assert all(block["pageNumber"] is None for block in body["blocks"])


def test_parse_pdf_keeps_real_one_based_page_number(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={"file": ("manual.pdf", _pdf_bytes(), "application/pdf")},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["fileType"] == "pdf"
    assert body["pageCount"] == 1
    assert body["blocks"][0]["pageNumber"] == 1
    assert "Hello PDF" in body["blocks"][0]["text"]


def test_parse_rejects_pdf_without_extractable_text(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={
            "file": (
                "scanned.pdf",
                _pdf_bytes(include_text=False),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 422
    assert response.json()["code"] == "NO_TEXT_EXTRACTED"


def test_parse_rejects_password_protected_pdf(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={
            "file": (
                "protected.pdf",
                _encrypted_pdf_bytes(),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 422
    assert response.json()["code"] == "PDF_PASSWORD_REQUIRED"


def test_parse_accepts_pdf_with_empty_user_password(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={
            "file": (
                "permission-restricted.pdf",
                _pdf_bytes(encryption_password=""),
                "application/pdf",
            )
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert body["pageCount"] == 1
    assert body["blocks"][0]["pageNumber"] == 1
    assert "Hello PDF" in body["blocks"][0]["text"]


@pytest.mark.parametrize("file_name", ["legacy.doc", "archive.exe"])
def test_parse_rejects_unsupported_extension(
    client: TestClient,
    auth_headers: dict[str, str],
    file_name: str,
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={"file": (file_name, b"content", "application/octet-stream")},
    )

    assert response.status_code == 415
    assert response.json()["code"] == "UNSUPPORTED_FILE_TYPE"


def test_parse_rejects_empty_file(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={"file": ("empty.txt", b"", "text/plain")},
    )

    assert response.status_code == 422
    assert response.json()["code"] == "FILE_EMPTY"


def test_parse_rejects_file_over_configured_limit(
    client: TestClient,
    auth_headers: dict[str, str],
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setenv("OFFICECARE_MAX_DOCUMENT_SIZE_MB", "1")
    get_settings.cache_clear()

    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={"file": ("large.txt", b"x" * (1024 * 1024 + 1), "text/plain")},
    )

    assert response.status_code == 413
    assert response.json()["code"] == "FILE_TOO_LARGE"
    assert response.json()["details"]["maxSizeBytes"] == 1024 * 1024


def test_parse_rejects_corrupt_pdf(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={"file": ("broken.pdf", b"%PDF-not-valid", "application/pdf")},
    )

    assert response.status_code == 422
    assert response.json()["code"] == "INVALID_PDF"


def test_parse_rejects_corrupt_docx(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    response = client.post(
        PARSE_URL,
        headers=auth_headers,
        files={
            "file": (
                "broken.docx",
                b"not-a-docx-package",
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 422
    assert response.json()["code"] == "INVALID_DOCX"


def test_trace_id_is_propagated_to_success_and_error(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    headers = {**auth_headers, "X-Trace-Id": "officecare-test-001"}
    success = client.post(
        PARSE_URL,
        headers=headers,
        files={"file": ("manual.txt", b"hello", "text/plain")},
    )
    error = client.post(
        PARSE_URL,
        headers={"X-Trace-Id": "officecare-test-002"},
        files={"file": ("manual.txt", b"hello", "text/plain")},
    )

    assert success.status_code == 200
    assert success.headers["X-Trace-Id"] == "officecare-test-001"
    assert error.status_code == 401
    assert error.headers["X-Trace-Id"] == "officecare-test-002"
    assert error.json()["traceId"] == "officecare-test-002"
